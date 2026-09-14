"""Exportación de registros: Excel, CSV y expediente para la Inspección.

La versión de 2024 exportaba con *pandas* dos hojas planas sin totales ni
identificación de la empresa.  Aquí se usa *openpyxl* directamente: evita
arrastrar pandas y NumPy (unos 60 MB de dependencias que complicaban la
instalación) y permite dar formato al informe.

El art. 34.9 ET obliga a que el registro esté a disposición de las personas
trabajadoras, de sus representantes legales y de la Inspección de Trabajo, así
que hay tres salidas distintas:

* ``exportar_excel``      → informe de empresa, legible por una persona.
* ``exportar_trabajador`` → copia del registro propio de un trabajador.
* ``expediente_itss``     → volcado íntegro y verificable en JSON.
"""

from __future__ import annotations

import csv
import datetime as dt
import json
import sqlite3
from pathlib import Path
from typing import Any
from collections.abc import Sequence

from . import db, normativa
from .config import Ajustes, directorio_exportaciones
from .dominio import (
    Jornada,
    Trabajador,
    eventos_efectivos,
    formatear_horas,
    jornadas_de,
)
from .seguridad import Cifrador
from .version import __version__


class ErrorInforme(Exception):
    """No se ha podido generar el informe."""


def _marca_tiempo() -> str:
    return dt.datetime.now().strftime("%Y%m%d_%H%M%S")


def _destino(nombre: str, carpeta: Path | None = None) -> Path:
    carpeta = carpeta or directorio_exportaciones()
    carpeta.mkdir(parents=True, exist_ok=True)
    return carpeta / nombre


def _fila_jornada(jornada: Jornada, trabajador: Trabajador, ajustes: Ajustes) -> list[Any]:
    inicio = db.a_local(jornada.inicio)
    fin = db.a_local(jornada.fin) if jornada.fin else None
    trabajados = jornada.minutos_trabajados(ajustes.pausas_computan_como_trabajo)
    return [
        trabajador.codigo,
        trabajador.nombre,
        inicio.strftime("%d/%m/%Y"),
        inicio.strftime("%H:%M"),
        fin.strftime("%H:%M") if fin else "— sin salida —",
        jornada.minutos_pausa,
        round(trabajados / 60, 2),
        jornada.modalidad.capitalize(),
        "Sí" if jornada.incidencia else "",
        "Sí" if jornada.rectificada else "",
        _procedencia(jornada),
        jornada.nota,
    ]


_CABECERA_JORNADAS = [
    "Código", "Trabajador", "Fecha", "Entrada", "Salida", "Pausa (min)",
    "Horas trabajadas", "Modalidad", "Incidencia", "Rectificada", "Procedencia",
    "Observaciones",
]


def _procedencia(jornada: Jornada) -> str:
    """De dónde sale el registro. Ante una inspección esto no es un detalle.

    Un fichaje hecho en el terminal se registró cuando ocurrió; uno importado
    del programa anterior entró en el sistema mucho después, y el informe debe
    decirlo en vez de presentarlos como equivalentes.
    """
    if not jornada.importada:
        return "Fichado en el terminal"
    if jornada.registrada is None:
        return "Importado del sistema anterior"
    return (
        "Importado del sistema anterior el "
        f"{db.a_local(jornada.registrada):%d/%m/%Y}"
    )


def _resumen_procedencia(conexion: sqlite3.Connection) -> dict[str, Any]:
    """Cuántos fichajes se registraron aquí y cuántos vienen del sistema viejo."""
    filas = {
        f["origen"]: (f["n"], f["creado"])
        for f in conexion.execute(
            "SELECT origen, COUNT(*) AS n, MIN(creado_utc) AS creado "
            "FROM eventos GROUP BY origen"
        )
    }
    importados = filas.get("MIGRADO", (0, None))
    nativos = sum(n for origen, (n, _) in filas.items() if origen != "MIGRADO")
    fecha = None
    if importados[1]:
        fecha = db.a_local(db.desde_iso(importados[1])).strftime("%d/%m/%Y")
    return {
        "nativos": nativos,
        "importados": importados[0],
        "fecha_importacion": fecha,
    }


def _texto_cortes(integridad: dict) -> list[str]:
    """Explica los huecos que ha dejado una purga por caducidad legal.

    Un hueco en la cadena no es lo mismo que una manipulación, pero desde
    fuera se parece: conviene que el informe lo diga antes de que lo pregunten.
    """
    cortes = integridad["eventos"].get("cortes") or []
    if not cortes:
        return []
    purgados = sum(c.get("purgados", 0) for c in cortes)
    fechas = sorted({(c.get("cuando") or "")[:10] for c in cortes if c.get("cuando")})
    lineas = [
        f"La cadena tiene {len(cortes)} corte(s) por la purga de registros ya "
        f"caducados: se retiraron {purgados} fichaje(s) que habían superado el "
        "plazo legal de conservación.",
        "No es una alteración: la purga está anotada en el registro de "
        "auditoría, con fecha y autor, y el resto de la cadena cuadra.",
    ]
    if fechas:
        lineas.append("Fecha(s) de purga: " + ", ".join(fechas) + ".")
    return lineas


def _texto_alcance(procedencia: dict[str, Any]) -> list[str]:
    """Explicación honesta de hasta dónde llega la garantía de integridad."""
    lineas = [
        "La cadena de huellas SHA-256 acredita que los registros no se han",
        "modificado desde que entraron en esta aplicación.",
    ]
    if procedencia["importados"]:
        fecha = procedencia["fecha_importacion"] or "la fecha de importación"
        lineas += [
            "",
            f"{procedencia['importados']} de estos fichajes proceden del programa",
            f"anterior y se incorporaron el {fecha}. Para ellos la verificación",
            "acredita únicamente que no se han alterado desde esa fecha: lo",
            "ocurrido antes depende del sistema anterior, no de este.",
            "La columna «Procedencia» de la hoja Jornadas indica el origen de",
            "cada registro.",
        ]
    else:
        lineas += [
            "",
            "Todos los fichajes se registraron directamente en el terminal, en el",
            "momento en que se produjeron.",
        ]
    return lineas


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #

def exportar_excel(
    conexion: sqlite3.Connection,
    cifrador: Cifrador,
    ajustes: Ajustes,
    trabajadores: Sequence[Trabajador],
    desde: dt.datetime,
    hasta: dt.datetime,
    *,
    carpeta: Path | None = None,
) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise ErrorInforme(
            "Falta la dependencia 'openpyxl' para exportar a Excel.\n"
            "Instálala con:  pip install openpyxl\n"
            "Mientras tanto puedes exportar a CSV."
        ) from exc

    libro = Workbook()
    azul = PatternFill("solid", fgColor="1F3A5F")
    blanco_negrita = Font(color="FFFFFF", bold=True)

    def escribir_cabecera(hoja, cabecera: list[str]) -> None:
        hoja.append(cabecera)
        for celda in hoja[1]:
            celda.fill = azul
            celda.font = blanco_negrita
            celda.alignment = Alignment(horizontal="center", vertical="center")
        hoja.freeze_panes = "A2"

    def ajustar(hoja) -> None:
        for columna in hoja.columns:
            ancho = max((len(str(c.value)) for c in columna if c.value), default=8)
            hoja.column_dimensions[
                get_column_letter(columna[0].column)
            ].width = min(max(ancho + 3, 11), 46)

    # -- Resumen ------------------------------------------------------------ #
    resumen = libro.active
    resumen.title = "Resumen"
    resumen["A1"] = "Registro de jornada"
    resumen["A1"].font = Font(size=16, bold=True, color="1F3A5F")
    filas_meta = [
        ("Empresa", ajustes.empresa or "(sin indicar)"),
        ("CIF/NIF", ajustes.cif or "(sin indicar)"),
        ("Centro de trabajo", ajustes.centro_trabajo or "(sin indicar)"),
        ("Periodo", f"{db.a_local(desde):%d/%m/%Y} a {db.a_local(hasta):%d/%m/%Y}"),
        ("Generado", dt.datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Aplicación", f"Control Horario {__version__}"),
        ("Base legal", "art. 34.9 del Estatuto de los Trabajadores"),
        ("Conservación", f"{ajustes.anios_conservacion} años"),
    ]
    for indice, (etiqueta, valor) in enumerate(filas_meta, start=3):
        resumen.cell(row=indice, column=1, value=etiqueta).font = Font(bold=True)
        resumen.cell(row=indice, column=2, value=valor)

    cabecera_totales = [
        "Código", "Trabajador", "Jornadas", "Días", "Horas trabajadas",
        "Pausas (h)", "Incidencias", "Rectificadas", "Sin cerrar",
    ]
    inicio_totales = len(filas_meta) + 5
    for columna, titulo in enumerate(cabecera_totales, start=1):
        celda = resumen.cell(row=inicio_totales, column=columna, value=titulo)
        celda.fill = azul
        celda.font = blanco_negrita

    total_minutos = 0
    for desplazamiento, trabajador in enumerate(trabajadores, start=1):
        jornadas = jornadas_de(conexion, trabajador.id, desde=desde, hasta=hasta)
        minutos = sum(
            j.minutos_trabajados(ajustes.pausas_computan_como_trabajo)
            for j in jornadas
        )
        total_minutos += minutos
        valores = [
            trabajador.codigo,
            trabajador.nombre,
            len(jornadas),
            len({j.fecha for j in jornadas}),
            round(minutos / 60, 2),
            round(sum(j.minutos_pausa for j in jornadas) / 60, 2),
            sum(1 for j in jornadas if j.incidencia),
            sum(1 for j in jornadas if j.rectificada),
            sum(1 for j in jornadas if j.abierta),
        ]
        for columna, valor in enumerate(valores, start=1):
            resumen.cell(row=inicio_totales + desplazamiento, column=columna, value=valor)

    fila_total = inicio_totales + len(trabajadores) + 1
    resumen.cell(row=fila_total, column=2, value="TOTAL").font = Font(bold=True)
    celda_total = resumen.cell(row=fila_total, column=5, value=round(total_minutos / 60, 2))
    celda_total.font = Font(bold=True)
    ajustar(resumen)

    # -- Jornadas ------------------------------------------------------------ #
    hoja_jornadas = libro.create_sheet("Jornadas")
    escribir_cabecera(hoja_jornadas, _CABECERA_JORNADAS)
    for trabajador in trabajadores:
        for jornada in jornadas_de(conexion, trabajador.id, desde=desde, hasta=hasta):
            hoja_jornadas.append(_fila_jornada(jornada, trabajador, ajustes))
    hoja_jornadas.auto_filter.ref = hoja_jornadas.dimensions
    ajustar(hoja_jornadas)

    # -- Fichajes ------------------------------------------------------------ #
    hoja_fichajes = libro.create_sheet("Fichajes")
    escribir_cabecera(
        hoja_fichajes,
        ["Código", "Trabajador", "Fecha y hora", "Tipo", "Modalidad", "Origen",
         "Autor", "Rectificado", "Hora original", "Motivo"],
    )
    for trabajador in trabajadores:
        for evento in eventos_efectivos(
            conexion, trabajador_id=trabajador.id, desde=desde, hasta=hasta
        ):
            hoja_fichajes.append(
                [
                    trabajador.codigo,
                    trabajador.nombre,
                    db.a_local(evento.momento).strftime("%d/%m/%Y %H:%M:%S"),
                    evento.tipo.replace("_", " ").capitalize(),
                    evento.modalidad.capitalize(),
                    evento.origen,
                    evento.autor,
                    "Sí" if evento.rectificado else "",
                    db.a_local(evento.momento_original).strftime("%d/%m/%Y %H:%M")
                    if evento.momento_original
                    else "",
                    evento.motivo,
                ]
            )
    hoja_fichajes.auto_filter.ref = hoja_fichajes.dimensions
    ajustar(hoja_fichajes)

    # -- Alertas -------------------------------------------------------------- #
    hoja_alertas = libro.create_sheet("Alertas")
    escribir_cabecera(
        hoja_alertas, ["Gravedad", "Código", "Fecha", "Trabajador", "Detalle", "Norma"]
    )
    alertas = normativa.analizar_periodo(conexion, trabajadores, ajustes, desde, hasta)
    rojo = PatternFill("solid", fgColor="F8D7DA")
    ambar = PatternFill("solid", fgColor="FFF3CD")
    for alerta in alertas:
        hoja_alertas.append(
            [
                alerta.gravedad.upper(),
                alerta.codigo,
                alerta.fecha_texto,
                alerta.trabajador,
                alerta.mensaje,
                alerta.referencia,
            ]
        )
        relleno = rojo if alerta.gravedad == normativa.GRAVE else ambar
        for celda in hoja_alertas[hoja_alertas.max_row]:
            celda.fill = relleno
    if not alertas:
        hoja_alertas.append(["—", "", "", "", "Sin alertas en el periodo.", ""])
    ajustar(hoja_alertas)

    # -- Integridad ----------------------------------------------------------- #
    hoja_integridad = libro.create_sheet("Integridad")
    informe = db.verificar_integridad(conexion)
    hoja_integridad["A1"] = "Verificación de integridad del registro"
    hoja_integridad["A1"].font = Font(size=13, bold=True, color="1F3A5F")
    filas = [
        ("Comprobado", informe["verificado_utc"]),
        ("Resultado", "ÍNTEGRO" if informe["integro"] else "ALTERADO"),
        ("Fichajes encadenados", informe["eventos"]["filas"]),
        ("Hash final de fichajes", informe["eventos"]["hash_final"]),
        ("Registros de auditoría", informe["auditoria"]["filas"]),
        ("Hash final de auditoría", informe["auditoria"]["hash_final"]),
    ]
    for indice, (etiqueta, valor) in enumerate(filas, start=3):
        hoja_integridad.cell(row=indice, column=1, value=etiqueta).font = Font(bold=True)
        hoja_integridad.cell(row=indice, column=2, value=str(valor))
    fila = len(filas) + 4

    # Alcance real de la verificación. Presentar como equivalentes un fichaje
    # hecho en el terminal y uno importado sería inducir a error.
    procedencia = _resumen_procedencia(conexion)
    hoja_integridad.cell(
        row=fila, column=1, value="Qué acredita esta verificación"
    ).font = Font(bold=True, size=11)
    fila += 1
    for linea in _texto_alcance(procedencia):
        hoja_integridad.cell(row=fila, column=1, value=linea)
        fila += 1
    for linea in _texto_cortes(informe):
        hoja_integridad.cell(row=fila, column=1, value=linea)
        fila += 1

    if procedencia["importados"]:
        fila += 1
        hoja_integridad.cell(
            row=fila, column=1, value="Procedencia de los fichajes"
        ).font = Font(bold=True, size=11)
        fila += 1
        for etiqueta, cuantos in (
            ("Fichados en el terminal", procedencia["nativos"]),
            ("Importados del sistema anterior", procedencia["importados"]),
        ):
            hoja_integridad.cell(row=fila, column=1, value=etiqueta)
            hoja_integridad.cell(row=fila, column=2, value=cuantos)
            fila += 1
        if procedencia["fecha_importacion"]:
            hoja_integridad.cell(row=fila, column=1, value="Fecha de la importación")
            hoja_integridad.cell(
                row=fila, column=2, value=procedencia["fecha_importacion"]
            )
            fila += 1

    fila += 1
    for incidencia in informe["eventos"]["incidencias"] + informe["auditoria"]["incidencias"]:
        hoja_integridad.cell(
            row=fila, column=1,
            value=f"Fichaje {incidencia['id']}: {incidencia['problema']}",
        ).fill = rojo
        hoja_integridad.cell(row=fila, column=2, value=incidencia["detalle"])
        fila += 1
    ajustar(hoja_integridad)

    ruta = _destino(f"registro_jornada_{_marca_tiempo()}.xlsx", carpeta)
    libro.save(ruta)
    return ruta


# --------------------------------------------------------------------------- #
# CSV
# --------------------------------------------------------------------------- #

def exportar_csv(
    conexion: sqlite3.Connection,
    ajustes: Ajustes,
    trabajadores: Sequence[Trabajador],
    desde: dt.datetime,
    hasta: dt.datetime,
    *,
    carpeta: Path | None = None,
) -> Path:
    """CSV con separador ';' y BOM, que es lo que espera Excel en español."""
    ruta = _destino(f"registro_jornada_{_marca_tiempo()}.csv", carpeta)
    with ruta.open("w", encoding="utf-8-sig", newline="") as fichero:
        escritor = csv.writer(fichero, delimiter=";")
        escritor.writerow(_CABECERA_JORNADAS)
        for trabajador in trabajadores:
            for jornada in jornadas_de(
                conexion, trabajador.id, desde=desde, hasta=hasta
            ):
                fila = _fila_jornada(jornada, trabajador, ajustes)
                fila[6] = str(fila[6]).replace(".", ",")  # decimal español
                escritor.writerow(fila)
    return ruta


# --------------------------------------------------------------------------- #
# Registro individual del trabajador
# --------------------------------------------------------------------------- #

def exportar_trabajador(
    conexion: sqlite3.Connection,
    ajustes: Ajustes,
    trabajador: Trabajador,
    desde: dt.datetime,
    hasta: dt.datetime,
    *,
    carpeta: Path | None = None,
) -> Path:
    """Copia del registro propio, en texto plano legible.

    El art. 34.9 ET reconoce a la persona trabajadora el derecho de acceso a su
    registro; entregarlo en un formato abierto evita discusiones sobre si pudo
    consultarlo.
    """
    jornadas = jornadas_de(conexion, trabajador.id, desde=desde, hasta=hasta)
    total = sum(
        j.minutos_trabajados(ajustes.pausas_computan_como_trabajo) for j in jornadas
    )

    lineas = [
        "REGISTRO DE JORNADA",
        "=" * 68,
        f"Empresa           : {ajustes.empresa or '(sin indicar)'}",
        f"Centro de trabajo : {ajustes.centro_trabajo or '(sin indicar)'}",
        f"Trabajador        : {trabajador.nombre} ({trabajador.codigo})",
        f"DNI               : {trabajador.dni or '(no consta)'}",
        f"Periodo           : {db.a_local(desde):%d/%m/%Y} a {db.a_local(hasta):%d/%m/%Y}",
        f"Expedido          : {dt.datetime.now():%d/%m/%Y %H:%M}",
        "",
        f"{'Fecha':<12}{'Entrada':<10}{'Salida':<10}{'Pausa':<8}{'Trabajado':<12}Obs.",
        "-" * 68,
    ]
    for jornada in jornadas:
        inicio = db.a_local(jornada.inicio)
        fin = db.a_local(jornada.fin) if jornada.fin else None
        observaciones = []
        if jornada.incidencia:
            observaciones.append("incidencia")
        if jornada.rectificada:
            observaciones.append("rectificada")
        if jornada.abierta:
            observaciones.append("sin salida")
        trabajado = formatear_horas(
            jornada.minutos_trabajados(ajustes.pausas_computan_como_trabajo)
        )
        lineas.append(
            f"{inicio:%d/%m/%Y}".ljust(12)
            + f"{inicio:%H:%M}".ljust(10)
            + (f"{fin:%H:%M}" if fin else "--:--").ljust(10)
            + f"{jornada.minutos_pausa} min".ljust(8)
            + trabajado.ljust(12)
            + ", ".join(observaciones)
        )

    lineas += [
        "-" * 68,
        f"TOTAL DEL PERIODO : {formatear_horas(total)} en {len(jornadas)} jornada(s)",
        "",
        "Este documento reproduce el registro diario de jornada previsto en el",
        "artículo 34.9 del Estatuto de los Trabajadores. Si detectas un error,",
        "solicita su rectificación: quedará anotada con fecha, autor y motivo,",
        "sin borrar el dato original.",
    ]

    nombre = trabajador.nombre.replace(" ", "_").replace("/", "-")
    ruta = _destino(f"registro_{trabajador.codigo}_{nombre}_{_marca_tiempo()}.txt", carpeta)
    ruta.write_text("\n".join(lineas), encoding="utf-8")
    return ruta


# --------------------------------------------------------------------------- #
# Expediente para la Inspección de Trabajo
# --------------------------------------------------------------------------- #

def expediente_itss(
    conexion: sqlite3.Connection,
    ajustes: Ajustes,
    trabajadores: Sequence[Trabajador],
    desde: dt.datetime,
    hasta: dt.datetime,
    *,
    carpeta: Path | None = None,
) -> Path:
    """Volcado íntegro en JSON, con la cadena de hashes para verificarlo.

    El borrador de real decreto en tramitación apunta a un acceso telemático y
    a un formato interoperable.  Mientras no se apruebe, un JSON documentado y
    verificable cubre el deber de puesta a disposición del art. 34.9 ET y deja
    el dato preparado para volcarlo a lo que finalmente se exija.
    """
    integridad = db.verificar_integridad(conexion)

    personas = []
    for trabajador in trabajadores:
        jornadas = jornadas_de(conexion, trabajador.id, desde=desde, hasta=hasta)
        personas.append(
            {
                "codigo": trabajador.codigo,
                "nombre": trabajador.nombre,
                "dni": trabajador.dni,
                "rol": trabajador.rol,
                "activo": trabajador.activo,
                "jornada_semanal_pactada": trabajador.jornada_semanal
                or ajustes.horas_semanales,
                "jornadas": [
                    {
                        "fecha": db.a_local(j.inicio).strftime("%Y-%m-%d"),
                        "entrada": db.a_iso(j.inicio),
                        "salida": db.a_iso(j.fin) if j.fin else None,
                        "modalidad": j.modalidad,
                        "pausas": [
                            {
                                "inicio": db.a_iso(p_ini),
                                "fin": db.a_iso(p_fin) if p_fin else None,
                            }
                            for p_ini, p_fin in j.pausas
                        ],
                        "minutos_pausa": j.minutos_pausa,
                        "minutos_trabajados": j.minutos_trabajados(
                            ajustes.pausas_computan_como_trabajo
                        ),
                        "incidencia": j.incidencia,
                        "rectificada": j.rectificada,
                        "origen": j.origen,
                        "registrada_utc": db.a_iso(j.registrada)
                        if j.registrada
                        else None,
                    }
                    for j in jornadas
                ],
                "fichajes": [
                    {
                        "id": e.id,
                        "tipo": e.tipo,
                        "momento_utc": db.a_iso(e.momento),
                        "modalidad": e.modalidad,
                        "origen": e.origen,
                        "autor": e.autor,
                        "rectificado": e.rectificado,
                        "momento_original_utc": db.a_iso(e.momento_original)
                        if e.momento_original
                        else None,
                        "motivo_rectificacion": e.motivo,
                    }
                    for e in eventos_efectivos(
                        conexion, trabajador_id=trabajador.id, desde=desde, hasta=hasta
                    )
                ],
            }
        )

    auditoria = [
        dict(fila)
        for fila in conexion.execute(
            "SELECT id, ts_utc, actor, accion, detalle, hash FROM auditoria "
            "WHERE ts_utc BETWEEN ? AND ? ORDER BY id",
            (db.a_iso(desde), db.a_iso(hasta)),
        )
    ]

    expediente = {
        "formato": "controlhorario/expediente-itss",
        "version_formato": 1,
        "generado_utc": db.a_iso(db.ahora_utc()),
        "aplicacion": f"Control Horario {__version__}",
        "empresa": {
            "razon_social": ajustes.empresa,
            "cif": ajustes.cif,
            "centro_trabajo": ajustes.centro_trabajo,
        },
        "marco_legal": {
            "obligacion": "art. 34.9 del Estatuto de los Trabajadores",
            "introducida_por": "Real Decreto-ley 8/2019, de 8 de marzo",
            "conservacion_anios": ajustes.anios_conservacion,
        },
        "periodo": {"desde_utc": db.a_iso(desde), "hasta_utc": db.a_iso(hasta)},
        "integridad": integridad,
        "procedencia": _resumen_procedencia(conexion),
        "alcance_de_la_verificacion": " ".join(
            _texto_alcance(_resumen_procedencia(conexion))
        ).replace("  ", " ").strip(),
        "trabajadores": personas,
        "auditoria": auditoria,
        "alertas": [
            {
                "codigo": a.codigo,
                "gravedad": a.gravedad,
                "trabajador": a.trabajador,
                "fecha": a.fecha_texto,
                "mensaje": a.mensaje,
                "referencia": a.referencia,
            }
            for a in normativa.analizar_periodo(
                conexion, trabajadores, ajustes, desde, hasta
            )
        ],
    }

    ruta = _destino(f"expediente_itss_{_marca_tiempo()}.json", carpeta)
    ruta.write_text(
        json.dumps(expediente, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    return ruta


__all__ = [
    "ErrorInforme",
    "expediente_itss",
    "exportar_csv",
    "exportar_excel",
    "exportar_trabajador",
]


# --------------------------------------------------------------------------- #
# Informe imprimible para la Inspección
# --------------------------------------------------------------------------- #

_ESTILO_IMPRESION = """
:root { --tinta:#111827; --suave:#6B7280; --linea:#D1D5DB; --marca:#1F3A5F; }
* { box-sizing:border-box; }
body { font-family:"Segoe UI",system-ui,-apple-system,"Helvetica Neue",sans-serif;
       color:var(--tinta); margin:0; padding:24px; font-size:12px; line-height:1.45; }
.aviso { background:#FEF3C7; border:1px solid #F59E0B; border-radius:6px;
         padding:12px 16px; margin-bottom:20px; font-size:12px; }
h1 { font-size:20px; margin:0 0 4px; color:var(--marca); }
h2 { font-size:14px; margin:26px 0 8px; color:var(--marca);
     border-bottom:2px solid var(--marca); padding-bottom:4px; }
.sub { color:var(--suave); margin:0 0 18px; }
table { width:100%; border-collapse:collapse; margin-bottom:6px; }
th { background:var(--marca); color:#fff; text-align:left; padding:6px 8px;
     font-size:11px; font-weight:600; }
td { padding:5px 8px; border-bottom:1px solid var(--linea); }
tr:nth-child(even) td { background:#F9FAFB; }
td.n, th.n { text-align:right; }
.ficha { border:1px solid var(--linea); border-radius:6px; padding:12px 16px;
         margin-bottom:16px; }
.ficha dl { display:grid; grid-template-columns:auto 1fr auto 1fr; gap:4px 14px;
            margin:0; }
.ficha dt { color:var(--suave); }
.ficha dd { margin:0; font-weight:600; }
.total { font-weight:700; background:#EFF6FF !important; }
.marca { color:#92400E; font-size:10px; }
.pie { margin-top:26px; padding-top:12px; border-top:1px solid var(--linea);
       color:var(--suave); font-size:11px; }
.firma { margin-top:34px; display:flex; gap:60px; }
.firma div { flex:1; border-top:1px solid var(--tinta); padding-top:6px;
             color:var(--suave); }
@media print {
  body { padding:0; font-size:10.5px; }
  .aviso { display:none; }
  h2 { page-break-after:avoid; }
  table { page-break-inside:auto; }
  tr { page-break-inside:avoid; }
  .persona { page-break-before:always; }
  .persona:first-of-type { page-break-before:auto; }
  @page { size:A4; margin:14mm 12mm; }
}
"""


def _escapar(texto: object) -> str:
    from html import escape

    return escape(str(texto if texto is not None else ""))


def informe_inspeccion(
    conexion: sqlite3.Connection,
    ajustes: Ajustes,
    trabajadores: Sequence[Trabajador],
    desde: dt.datetime,
    hasta: dt.datetime,
    *,
    carpeta: Path | None = None,
) -> Path:
    """Documento legible y listo para imprimir o guardar en PDF.

    El Excel sirve para trabajar con los datos y el JSON para volcarlos, pero
    lo que se entrega en mano tiene que poder leerse sin abrir un programa de
    hojas de cálculo.  Se genera como página web autocontenida: se abre en
    cualquier navegador y con Ctrl+P se guarda en PDF, sin depender de más
    programas instalados.
    """
    procedencia = _resumen_procedencia(conexion)
    partes: list[str] = []

    partes.append(
        '<div class="aviso"><b>Para guardarlo en PDF:</b> pulsa '
        "<b>Ctrl+P</b> y elige «Guardar como PDF» o «Microsoft Print to PDF». "
        "Este aviso no sale impreso.</div>"
    )
    partes.append("<h1>Registro de jornada</h1>")
    partes.append(
        '<p class="sub">Artículo 34.9 del Estatuto de los Trabajadores</p>'
    )

    partes.append('<div class="ficha"><dl>')
    for etiqueta, valor in (
        ("Empresa", ajustes.empresa or "—"),
        ("CIF / NIF", ajustes.cif or "—"),
        ("Centro de trabajo", ajustes.centro_trabajo or "—"),
        ("Periodo", f"{db.a_local(desde):%d/%m/%Y} a {db.a_local(hasta):%d/%m/%Y}"),
        ("Expedido", dt.datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Personas incluidas", str(len(trabajadores))),
    ):
        partes.append(f"<dt>{_escapar(etiqueta)}</dt><dd>{_escapar(valor)}</dd>")
    partes.append("</dl></div>")

    total_general = 0
    for trabajador in trabajadores:
        jornadas = jornadas_de(conexion, trabajador.id, desde=desde, hasta=hasta)
        if not jornadas:
            continue
        partes.append('<div class="persona">')
        partes.append(
            f"<h2>{_escapar(trabajador.nombre)} · {_escapar(trabajador.codigo)}"
            + (f" · DNI {_escapar(trabajador.dni)}" if trabajador.dni else "")
            + "</h2>"
        )
        # Si toda la serie viene del programa anterior, se dice una vez bajo el
        # nombre en lugar de repetirlo en cada fila: se lee mucho mejor y el
        # dato queda igual de claro.
        importadas = sum(1 for j in jornadas if j.importada)
        todas_importadas = importadas == len(jornadas)
        if todas_importadas:
            fecha = next(
                (db.a_local(j.registrada).strftime("%d/%m/%Y")
                 for j in jornadas if j.registrada),
                None,
            )
            partes.append(
                '<p class="marca">Todas las jornadas de este periodo proceden '
                "del programa anterior"
                + (f", incorporadas el {fecha}" if fecha else "")
                + ".</p>"
            )
        partes.append(
            "<table><thead><tr>"
            "<th>Fecha</th><th>Entrada</th><th>Salida</th>"
            '<th class="n">Pausa</th><th class="n">Trabajado</th>'
            "<th>Modalidad</th><th>Observaciones</th>"
            "</tr></thead><tbody>"
        )
        total = 0
        for jornada in jornadas:
            inicio = db.a_local(jornada.inicio)
            fin = db.a_local(jornada.fin) if jornada.fin else None
            minutos = jornada.minutos_trabajados(
                ajustes.pausas_computan_como_trabajo
            )
            total += minutos
            observaciones = []
            if jornada.incidencia:
                observaciones.append("incidencia")
            if jornada.rectificada:
                observaciones.append("rectificada")
            if jornada.abierta:
                observaciones.append("sin fichaje de salida")
            if jornada.importada and not todas_importadas:
                observaciones.append(
                    '<span class="marca">importado del sistema anterior</span>'
                )
            salida = f"{fin:%H:%M}" if fin else "—"
            partes.append(
                "<tr>"
                f"<td>{inicio:%d/%m/%Y}</td>"
                f"<td>{inicio:%H:%M}</td>"
                f"<td>{salida}</td>"
                f'<td class="n">{jornada.minutos_pausa} min</td>'
                f'<td class="n">{formatear_horas(minutos)}</td>'
                f"<td>{_escapar(jornada.modalidad.capitalize())}</td>"
                f"<td>{' · '.join(observaciones)}</td>"
                "</tr>"
            )
        total_general += total
        partes.append(
            f'<tr class="total"><td colspan="4">TOTAL · {len(jornadas)} jornada(s)'
            f'</td><td class="n">{formatear_horas(total)}</td><td colspan="2"></td>'
            "</tr>"
        )
        partes.append("</tbody></table></div>")

    partes.append("<h2>Verificación de integridad</h2>")
    integridad = db.verificar_integridad(conexion)
    partes.append(
        "<p><b>Resultado: "
        + ("ÍNTEGRO" if integridad["integro"] else "ALTERADO")
        + f"</b> · {integridad['eventos']['filas']} fichajes encadenados · "
        f"huella final <code>{integridad['eventos']['hash_final'][:24]}…</code></p>"
    )
    partes.append("<p>" + " ".join(_texto_alcance(procedencia)) + "</p>")
    cortes = _texto_cortes(integridad)
    if cortes:
        partes.append("<p>" + _escapar(" ".join(cortes)) + "</p>")

    partes.append(
        '<div class="pie">'
        f"Total del periodo: <b>{formatear_horas(total_general)}</b>. "
        "El registro se conserva cuatro años y está a disposición de las "
        "personas trabajadoras, de sus representantes legales y de la "
        "Inspección de Trabajo y Seguridad Social."
        "</div>"
    )
    partes.append(
        '<div class="firma"><div>Firma y sello de la empresa</div>'
        "<div>Fecha</div></div>"
    )

    documento = (
        "<!doctype html><html lang='es'><head><meta charset='utf-8'>"
        f"<title>Registro de jornada · {_escapar(ajustes.empresa or 'empresa')}</title>"
        f"<style>{_ESTILO_IMPRESION}</style></head><body>"
        + "".join(partes)
        + "</body></html>"
    )

    ruta = _destino(f"informe_inspeccion_{_marca_tiempo()}.html", carpeta)
    ruta.write_text(documento, encoding="utf-8")
    return ruta


__all__ += ["informe_inspeccion"]
